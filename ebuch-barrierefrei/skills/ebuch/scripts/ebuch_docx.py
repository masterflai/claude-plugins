# -*- coding: utf-8 -*-
"""Baut Word-Dokumente nach E-Buch-Standard.

Kapselt die Fallstricke aus references/docx-fallstricke.md: Formatvorlagen für
deutsches Word, Aufzählungszeichen an den Liste-Vorlagen, Neustart nummerierter
Listen, schema-konforme Elementreihenfolge, Sprache, Inhaltsverzeichnis und das
Entfernen jeder Werkzeugspur aus den Metadaten.

    from ebuch_docx import EBuchDokument

    d = EBuchDokument(titel="Titel des Originals", autor="Urheber des Originals")
    d.vorbemerkung(quelle="PowerPoint-Präsentation", einheiten=40,
                   einheit="Folie", bilder_gesamt=20, bilder_verschieden=15,
                   dekorativ=["Logo auf jeder Folie"],
                   abweichungen=["Wiederholte Folientitel zusammengefasst."])
    d.inhaltsverzeichnis()
    d.ueberschrift("Teil 1", 1)
    d.marke(2); d.absatz("…"); d.liste(["a", "b"]); d.bild("…")
    d.speichern("ergebnis.docx")
"""
import re
import zipfile
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# --------------------------------------------------------------- Schema-Reihenfolgen

PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
    "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
    "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
]

RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
    "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish",
    "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight",
    "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
    "eastAsianLayout", "specVanish", "oMath",
]

SETTINGS_ORDER = [
    "writeProtection", "view", "zoom", "removePersonalInformation", "removeDateAndTime",
    "doNotDisplayPageBoundaries", "displayBackgroundShape", "printPostScriptOverText",
    "printFractionalCharacterWidth", "printFormsData", "embedTrueTypeFonts",
    "embedSystemFonts", "saveSubsetFonts", "saveFormsData", "mirrorMargins",
    "alignBordersAndEdges", "bordersDoNotSurroundHeader", "bordersDoNotSurroundFooter",
    "gutterAtTop", "hideSpellingErrors", "hideGrammaticalErrors", "activeWritingStyle",
    "proofState", "formsDesign", "attachedTemplate", "linkStyles", "stylePaneFormatFilter",
    "stylePaneSortMethod", "documentType", "mailMerge", "revisionView", "trackChanges",
    "doNotTrackMoves", "doNotTrackFormatting", "documentProtection", "autoFormatOverride",
    "styleLockTheme", "styleLockQFSet", "defaultTabStop", "autoHyphenation",
    "consecutiveHyphenLimit", "hyphenationZone", "doNotHyphenateCaps", "showEnvelope",
    "summaryLength", "clickAndTypeStyle", "defaultTableStyle", "evenAndOddHeaders",
    "bookFoldRevPrinting", "bookFoldPrinting", "bookFoldPrintingSheets",
    "drawingGridHorizontalSpacing", "drawingGridVerticalSpacing",
    "displayHorizontalDrawingGridEvery", "displayVerticalDrawingGridEvery",
    "doNotUseMarginsForDrawingGridOrigin", "drawingGridHorizontalOrigin",
    "drawingGridVerticalOrigin", "doNotShadeFormData", "noPunctuationKerning",
    "characterSpacingControl", "printTwoOnOne", "strictFirstAndLastChars",
    "noLineBreaksAfter", "noLineBreaksBefore", "savePreviewPicture",
    "doNotValidateAgainstSchema", "saveInvalidXml", "ignoreMixedContent",
    "alwaysShowPlaceholderText", "doNotDemarcateInvalidXml", "saveXmlDataOnly",
    "useXSLTWhenSaving", "saveThroughXslt", "showXMLTags", "alwaysMergeEmptyNamespace",
    "updateFields", "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars",
    "rsids", "mathPr", "uiCompat97To2003", "attachedSchema", "themeFontLang",
    "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures",
    "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
    "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator",
]

NEUTRALES_APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
    'extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/'
    '2006/docPropsVTypes"><Template/><TotalTime>0</TotalTime><Application/>'
    '<DocSecurity>0</DocSecurity><ScaleCrop>false</ScaleCrop><Company/>'
    '<LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc>'
    '<HyperlinksChanged>false</HyperlinksChanged><AppVersion>1.0000</AppVersion>'
    '</Properties>'
)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def _sortieren(eltern, reihenfolge):
    if eltern is None:
        return
    idx = {n: i for i, n in enumerate(reihenfolge)}
    kinder = list(eltern)
    if any(not isinstance(k.tag, str) for k in kinder):
        return
    kinder.sort(key=lambda k: idx.get(k.tag.split("}")[-1], len(reihenfolge) + 1))
    for k in kinder:
        eltern.append(k)


class _Nummerierung:
    """Eigene Listendefinitionen, damit Aufzählungen und Neustarts funktionieren."""

    def __init__(self, dokument):
        self.teil = dokument.part.numbering_part.element
        ids = [int(a.get(qn("w:abstractNumId")))
               for a in self.teil.findall(qn("w:abstractNum"))]
        self._naechste_abstrakt = max(ids) + 1 if ids else 0
        nums = [int(a.get(qn("w:numId"))) for a in self.teil.findall(qn("w:num"))]
        self._naechste_num = max(nums) + 1 if nums else 1

        self.abs_punkt = self._abstrakt([("bullet", "", 360), ("bullet", "o", 720)])
        self.abs_punkt2 = self._abstrakt([("bullet", "o", 720)])
        self.abs_zahl = self._abstrakt([("decimal", "%1.", 360)])
        self.num_punkt = self._num(self.abs_punkt)
        self.num_punkt2 = self._num(self.abs_punkt2)

    def _abstrakt(self, ebenen):
        a = _el("w:abstractNum", **{"w:abstractNumId": self._naechste_abstrakt})
        self._naechste_abstrakt += 1
        a.append(_el("w:multiLevelType", **{"w:val": "hybridMultilevel"}))
        for i, (fmt, zeichen, einzug) in enumerate(ebenen):
            lvl = _el("w:lvl", **{"w:ilvl": i})
            lvl.append(_el("w:start", **{"w:val": "1"}))
            lvl.append(_el("w:numFmt", **{"w:val": fmt}))
            lvl.append(_el("w:lvlText", **{"w:val": zeichen}))
            lvl.append(_el("w:lvlJc", **{"w:val": "left"}))
            ppr = OxmlElement("w:pPr")
            ppr.append(_el("w:ind", **{"w:left": str(einzug), "w:hanging": "360"}))
            lvl.append(ppr)
            if fmt == "bullet":
                rpr = OxmlElement("w:rPr")
                rpr.append(_el("w:rFonts", **{"w:ascii": "Symbol", "w:hAnsi": "Symbol",
                                              "w:hint": "default"}))
                lvl.append(rpr)
            a.append(lvl)
        vorhandene = self.teil.findall(qn("w:abstractNum"))
        if vorhandene:
            vorhandene[-1].addnext(a)
        else:
            self.teil.insert(0, a)
        return int(a.get(qn("w:abstractNumId")))

    def _num(self, abstrakt_id, neustart=False):
        n = _el("w:num", **{"w:numId": self._naechste_num})
        self._naechste_num += 1
        n.append(_el("w:abstractNumId", **{"w:val": str(abstrakt_id)}))
        if neustart:
            ov = _el("w:lvlOverride", **{"w:ilvl": "0"})
            ov.append(_el("w:startOverride", **{"w:val": "1"}))
            n.append(ov)
        self.teil.append(n)
        return int(n.get(qn("w:numId")))

    def neue_zahlenliste(self):
        return self._num(self.abs_zahl, neustart=True)


class EBuchDokument:
    """Word-Dokument nach E-Buch-Standard."""

    def __init__(self, titel, autor="", beschreibung="", schlagworte="",
                 schrift="Arial", groesse=12, marke_prefix="Folie"):
        self.titel_text = titel
        self.marke_prefix = marke_prefix
        self.doc = Document()
        self._seite()
        self._formatvorlagen(schrift, groesse)
        self.num = _Nummerierung(self.doc)
        self._listenvorlagen(schrift, groesse)
        self._sprache()
        self._einstellungen()
        self._eigenschaften(titel, autor, beschreibung, schlagworte)
        self.absatz_titel(titel)

    # ---------------------------------------------------------------- Aufbau

    def _seite(self):
        s = self.doc.sections[0]
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.0)

    def _farbe_schwarz(self, vorlage):
        rpr = vorlage.element.get_or_add_rPr()
        for alt in rpr.findall(qn("w:color")):
            rpr.remove(alt)
        rpr.append(_el("w:color", **{"w:val": "000000"}))

    def _kursiv_aus(self, vorlage):
        rpr = vorlage.element.get_or_add_rPr()
        for tag in ("w:i", "w:iCs"):
            for alt in rpr.findall(qn(tag)):
                rpr.remove(alt)
        rpr.append(_el("w:i", **{"w:val": "0"}))
        rpr.append(_el("w:iCs", **{"w:val": "0"}))

    def _formatvorlagen(self, schrift, groesse):
        vorhandene = [v.name for v in self.doc.styles]

        normal = self.doc.styles["Normal"]
        normal.font.name = schrift
        normal.font.size = Pt(groesse)
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.space_before = Pt(0)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), schrift)
        normal.element.rPr.rFonts.set(qn("w:cs"), schrift)

        for stufe, gr in ((1, groesse + 6), (2, groesse + 4),
                          (3, groesse + 2), (4, groesse)):
            v = self.doc.styles[f"Heading {stufe}"]
            v.font.name = schrift
            v.font.size = Pt(gr)
            v.font.bold = True
            self._kursiv_aus(v)
            self._farbe_schwarz(v)
            v.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            v.paragraph_format.space_before = Pt(14 if stufe <= 2 else 10)
            v.paragraph_format.space_after = Pt(6)
            v.paragraph_format.keep_with_next = True
            zeichen = f"Heading {stufe} Char"
            if zeichen in vorhandene:
                self._kursiv_aus(self.doc.styles[zeichen])
                self._farbe_schwarz(self.doc.styles[zeichen])

        t = self.doc.styles["Title"]
        t.font.name = schrift
        t.font.size = Pt(groesse + 12)
        t.font.bold = True
        self._farbe_schwarz(t)
        self._kursiv_aus(t)
        t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if "Title Char" in vorhandene:
            self._farbe_schwarz(self.doc.styles["Title Char"])
            self._kursiv_aus(self.doc.styles["Title Char"])

    def _listenvorlagen(self, schrift, groesse):
        def anhaengen(vorlage, num_id, einzug):
            ppr = vorlage.element.get_or_add_pPr()
            for alt in ppr.findall(qn("w:numPr")):
                ppr.remove(alt)
            npr = OxmlElement("w:numPr")
            npr.append(_el("w:ilvl", **{"w:val": "0"}))
            npr.append(_el("w:numId", **{"w:val": str(num_id)}))
            ppr.insert(0, npr)
            for alt in ppr.findall(qn("w:ind")):
                ppr.remove(alt)
            ppr.append(_el("w:ind", **{"w:left": str(einzug), "w:hanging": "360"}))

        anhaengen(self.doc.styles["List"], self.num.num_punkt, 360)
        anhaengen(self.doc.styles["List 2"], self.num.num_punkt2, 720)
        for name in ("List", "List 2", "List Number"):
            v = self.doc.styles[name]
            v.font.name = schrift
            v.font.size = Pt(groesse)
            v.paragraph_format.space_after = Pt(3)

    def _sprache(self, code="de-DE"):
        dd = self.doc.styles.element.find(qn("w:docDefaults"))
        rpr = dd.find(qn("w:rPrDefault")).find(qn("w:rPr"))
        for alt in rpr.findall(qn("w:lang")):
            rpr.remove(alt)
        rpr.append(_el("w:lang", **{"w:val": code, "w:eastAsia": code, "w:bidi": "ar-SA"}))
        self._sprachcode = code

    def _einstellungen(self):
        s = self.doc.settings.element
        zoom = s.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")
        s.append(_el("w:autoHyphenation", **{"w:val": "true"}))
        s.append(_el("w:updateFields", **{"w:val": "true"}))
        _sortieren(s, SETTINGS_ORDER)

    def _eigenschaften(self, titel, autor, beschreibung, schlagworte):
        cp = self.doc.core_properties
        cp.title = titel
        cp.author = autor or ""
        cp.last_modified_by = ""
        cp.comments = beschreibung or ""
        cp.keywords = schlagworte or ""
        cp.language = "de-DE"
        cp.category = ""

    # ---------------------------------------------------------------- Bausteine

    def absatz_titel(self, text):
        p = self.doc.add_paragraph(style="Title")
        p.add_run(text)
        return p

    def ueberschrift(self, text, ebene):
        p = self.doc.add_paragraph(style=f"Heading {ebene}")
        p.add_run(text)
        return p

    def absatz(self, text, vorlage=None):
        p = self.doc.add_paragraph(style=vorlage)
        p.add_run(text)
        return p

    def marke(self, nummer):
        """Seiten- bzw. Folienmarke: ((12)) oder ((Folie 12))."""
        kennung = f"{self.marke_prefix} {nummer}".strip() if self.marke_prefix else str(nummer)
        return self.absatz(f"(({kennung}))")

    def liste(self, eintraege, ebene=1):
        vorlage = "List" if ebene == 1 else "List 2"
        return [self.absatz(t, vorlage) for t in eintraege]

    def nummeriert(self, eintraege):
        """Nummerierte Liste, die garantiert bei 1 beginnt."""
        num_id = self.num.neue_zahlenliste()
        raus = []
        for t in eintraege:
            p = self.doc.add_paragraph(style="List Number")
            ppr = p._p.get_or_add_pPr()
            npr = OxmlElement("w:numPr")
            npr.append(_el("w:ilvl", **{"w:val": "0"}))
            npr.append(_el("w:numId", **{"w:val": str(num_id)}))
            ppr.insert(0, npr)
            p.add_run(t)
            raus.append(p)
        return raus

    def bild(self, beschreibung):
        self.absatz("<Bild>")
        self.absatz(beschreibung)
        self.absatz("</Bild>")

    def anmerkung(self, text):
        """Anmerkung der Übertragung, in Rauten."""
        return self.absatz(f"#{text}#")

    def luecke(self):
        return "_..._"

    def tabelle(self, kopf, zeilen, breiten=None):
        spalten, reihen = len(kopf), len(zeilen) + 1
        self.absatz(f"<Tabelle> Sp. {spalten}, Zei. {reihen}")
        t = self.doc.add_table(rows=1, cols=spalten)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = breiten is None
        kopfzeile = t.rows[0]
        kopfzeile._tr.get_or_add_trPr().append(_el("w:tblHeader", **{"w:val": "true"}))
        for i, text in enumerate(kopf):
            zelle = kopfzeile.cells[i]
            zelle.text = ""
            zelle.paragraphs[0].add_run(text).bold = True
        for zeile in zeilen:
            zellen = t.add_row().cells
            for i, wert in enumerate(zeile):
                zellen[i].text = ""
                zellen[i].paragraphs[0].add_run(wert)
        if breiten:
            for spalte, b in zip(t.columns, breiten):
                spalte.width = Cm(b)
            for reihe in t.rows:
                for zelle, b in zip(reihe.cells, breiten):
                    zelle.width = Cm(b)
        self.doc.add_paragraph()
        self.absatz("</Tabelle>")
        return t

    def inhaltsverzeichnis(self, ebenen="1-4",
                           platzhalter="Das Inhaltsverzeichnis wird beim Öffnen in Word "
                                       "aktualisiert. Falls nicht: mit F9 aktualisieren."):
        p = self.doc.add_paragraph()
        r = p.add_run()
        r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
        anweisung = OxmlElement("w:instrText")
        anweisung.set(qn("xml:space"), "preserve")
        anweisung.text = f' TOC \\o "{ebenen}" \\h \\z \\u '
        r._r.append(anweisung)
        r._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
        p.add_run(platzhalter)
        ende = p.add_run()
        ende._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
        return p

    def vorbemerkung(self, quelle, einheiten, einheit="Folie", bilder_gesamt=0,
                     bilder_verschieden=0, dekorativ=(), abweichungen=(),
                     bilder_eingebettet=False, fassung="Fassung 2025"):
        """Standardtext „Hinweise zu dieser Übertragung"."""
        self.ueberschrift("Hinweise zu dieser Übertragung", 1)
        self.absatz(f"Diese Datei ist die barrierefreie Textfassung der {quelle} "
                    f"„{self.titel_text}“. Sie wurde nach dem E-Buch-Standard übertragen "
                    f"(Handreichung in der aktualisierten {fassung}, "
                    f"herausgegeben auf augenbit.de).")

        self.ueberschrift("Aufbau", 2)
        self.absatz(f"Der Inhalt der {einheiten} {einheit}n ist fortlaufend als Text "
                    "gegliedert. Die Überschriften sind durchgängig mit den Formatvorlagen "
                    "Überschrift 1 bis Überschrift 4 ausgezeichnet, sodass das Dokument "
                    "über die Überschriftenliste des Screenreaders und über das "
                    "Inhaltsverzeichnis angesprungen werden kann.")

        self.ueberschrift("Verwendete Kennzeichnungen", 2)
        if self.marke_prefix:
            self.absatz(f"Vor jedem Inhalt steht die Nummer der Original{einheit.lower()} "
                        f"in doppelten runden Klammern, zum Beispiel (({self.marke_prefix} "
                        "17)). Der E-Buch-Standard verwendet diese Marke für Seitenzahlen; "
                        f"hier bezeichnet sie die {einheit}nnummer, damit im Unterricht auf "
                        "dieselbe Stelle Bezug genommen werden kann.", "List")
        else:
            self.absatz("Vor jedem Seiteninhalt steht die Seitenzahl des Originals in "
                        "doppelten runden Klammern, zum Beispiel ((17)).", "List")
        self.absatz("Bildbeschreibungen stehen zwischen den Marken <Bild> und </Bild>.",
                    "List")
        self.absatz("Tabellen stehen zwischen den Marken <Tabelle> und </Tabelle>. "
                    "Die Marke nennt zuerst die Anzahl der Spalten und Zeilen.", "List")
        self.absatz("Anmerkungen der Übertragung stehen zwischen Rauten, zum Beispiel "
                    "#Anmerkung#. Sie stammen nicht aus dem Original.", "List")
        self.absatz("Lücken, die im Unterricht ausgefüllt werden, sind als _..._ "
                    "gekennzeichnet.", "List")

        if bilder_gesamt:
            self.ueberschrift("Bilder", 2)
            satz = (f"Das Original enthält {bilder_gesamt} Grafiken, davon "
                    f"{bilder_verschieden} verschiedene. Alle inhaltlich bedeutsamen "
                    "Bilder sind beschrieben.")
            if dekorativ:
                satz += (f" Rein dekorative Grafiken sind nicht beschrieben: "
                         f"{'; '.join(dekorativ)}. Sie tragen keine eigene Information.")
            self.absatz(satz)
            if not bilder_eingebettet:
                self.absatz("Die Bilddateien selbst sind in dieser Fassung nicht "
                            "eingebettet, damit das Dokument schlank bleibt und sich "
                            "zuverlässig in Braille ausgeben lässt.")

        if abweichungen:
            self.ueberschrift("Abweichungen vom Original", 2)
            for a in abweichungen:
                self.absatz(a, "List")

    # ---------------------------------------------------------------- Speichern

    def speichern(self, pfad):
        for teil in (self.doc.element.body, self.doc.styles.element):
            for ppr in teil.iter(qn("w:pPr")):
                _sortieren(ppr, PPR_ORDER)
        for rpr in self.doc.styles.element.iter(qn("w:rPr")):
            _sortieren(rpr, RPR_ORDER)
        self.doc.save(pfad)
        _spuren_entfernen(pfad)
        return pfad


def _spuren_entfernen(pfad):
    """Neutralisiert docProps/app.xml und lastModifiedBy.

    In der erzeugten Datei darf nirgends stehen, womit sie gebaut wurde.
    """
    with zipfile.ZipFile(pfad) as z:
        eintraege = [(i, z.read(i.filename)) for i in z.infolist()]

    with zipfile.ZipFile(pfad, "w", zipfile.ZIP_DEFLATED) as z:
        for info, daten in eintraege:
            if info.filename == "docProps/app.xml":
                daten = NEUTRALES_APP_XML.encode("utf-8")
            elif info.filename == "docProps/core.xml":
                text = daten.decode("utf-8")
                text = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>",
                              "<cp:lastModifiedBy></cp:lastModifiedBy>", text, flags=re.S)
                daten = text.encode("utf-8")
            z.writestr(info, daten)
