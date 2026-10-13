#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Extrahiert Inhalt und Bilder aus Unterrichtsmaterial für die E-Buch-Übertragung.

    python extrahieren.py <quelldatei> --out <arbeitsordner>

Unterstützt .pptx (vollständig), .docx und .pdf (Text und Bilder).

Erzeugt im Arbeitsordner:
    inhalt.json   Struktur je Folie/Seite mit Lesereihenfolge
    bilder/       entdoppelte Grafiken, benannt nach ihrem SHA1-Präfix

Gibt auf der Standardausgabe eine indizierte Absatzliste aus. Die Indizes sind
die Bezugsgröße für den Strukturplan.
"""
import argparse
import json
import os
import re
import sys
from collections import Counter

AUTO_ALT = re.compile(r"Beschreibung automatisch generiert|automatically generated", re.I)


def saeubern(t):
    t = t.replace(" ", " ").replace("​", "").replace("‑", "-")
    return re.sub(r"[ \t]+", " ", t).strip()


def brauchbarer_alt(alt):
    return bool(alt and alt.strip() and not AUTO_ALT.search(alt))


# --------------------------------------------------------------------- PPTX

def aus_pptx(pfad, bilderordner):
    from pptx import Presentation

    praes = Presentation(pfad)
    einheiten = []
    bilder = {}

    for nr, folie in enumerate(praes.slides, 1):
        formen = []
        for form in folie.shapes:
            eintrag = {
                "name": form.name,
                "typ": str(form.shape_type),
                "top": form.top or 0,
                "left": form.left or 0,
                "breite": form.width or 0,
                "hoehe": form.height or 0,
                "platzhalter": None,
                "alt": "",
            }
            try:
                eintrag["platzhalter"] = str(form.placeholder_format.type)
            except Exception:
                pass
            try:
                eintrag["alt"] = form._element._nvXxPr.cNvPr.get("descr") or ""
            except Exception:
                pass

            try:
                bild = form.image
            except Exception:
                bild = None
            if bild is not None:
                sha = bild.sha1[:8]
                eintrag["bild"] = sha
                if sha not in bilder:
                    ziel = os.path.join(bilderordner, f"{sha}.{bild.ext}")
                    with open(ziel, "wb") as f:
                        f.write(bild.blob)
                    bilder[sha] = {"datei": os.path.basename(ziel),
                                   "pixel": list(bild.size), "folien": []}
                bilder[sha]["folien"].append(nr)

            if form.has_text_frame:
                eintrag["absaetze"] = [
                    {"ebene": p.level, "text": saeubern("".join(r.text for r in p.runs))}
                    for p in form.text_frame.paragraphs
                    if "".join(r.text for r in p.runs).strip()
                ]
            if form.has_table:
                eintrag["tabelle"] = [[saeubern(z.text) for z in reihe.cells]
                                      for reihe in form.table.rows]
            formen.append(eintrag)

        # Lesereihenfolge: Titel zuerst, dann von oben nach unten, links nach rechts
        def sortierung(e):
            ist_titel = 0 if (e["platzhalter"] and "TITLE" in e["platzhalter"]) else 1
            return (ist_titel, e["top"], e["left"])

        formen.sort(key=sortierung)
        notiz = ""
        if folie.has_notes_slide:
            notiz = saeubern(folie.notes_slide.notes_text_frame.text)
        einheiten.append({"nr": nr, "layout": folie.slide_layout.name,
                          "formen": formen, "notiz": notiz})
    return einheiten, bilder


# --------------------------------------------------------------------- DOCX

def aus_docx(pfad, bilderordner):
    import hashlib
    from docx import Document

    dok = Document(pfad)
    formen = []
    for p in dok.paragraphs:
        if p.text.strip():
            formen.append({"name": p.style.name, "typ": "TEXT", "top": 0, "left": 0,
                           "platzhalter": None, "alt": "",
                           "absaetze": [{"ebene": 0, "text": saeubern(p.text)}]})
    for t in dok.tables:
        formen.append({"name": "Tabelle", "typ": "TABLE", "top": 0, "left": 0,
                       "platzhalter": None, "alt": "",
                       "tabelle": [[saeubern(z.text) for z in r.cells] for r in t.rows]})

    bilder = {}
    for rel in dok.part.rels.values():
        if "image" in rel.reltype:
            blob = rel.target_part.blob
            sha = hashlib.sha1(blob).hexdigest()[:8]
            if sha not in bilder:
                endung = rel.target_part.partname.ext
                ziel = os.path.join(bilderordner, f"{sha}.{endung}")
                with open(ziel, "wb") as f:
                    f.write(blob)
                bilder[sha] = {"datei": os.path.basename(ziel), "pixel": None, "folien": [1]}
    return [{"nr": 1, "layout": "docx", "formen": formen, "notiz": ""}], bilder


# --------------------------------------------------------------------- PDF

def aus_pdf(pfad, bilderordner):
    import hashlib
    try:
        import fitz  # PyMuPDF
    except ImportError:
        sys.exit("Für PDF wird PyMuPDF benötigt: pip install pymupdf --break-system-packages")

    dok = fitz.open(pfad)
    einheiten, bilder = [], {}
    for nr, seite in enumerate(dok, 1):
        absaetze = [{"ebene": 0, "text": saeubern(b)}
                    for b in seite.get_text().split("\n") if b.strip()]
        formen = [{"name": "Seitentext", "typ": "TEXT", "top": 0, "left": 0,
                   "platzhalter": None, "alt": "", "absaetze": absaetze}]
        for bild_info in seite.get_images(full=True):
            blob = dok.extract_image(bild_info[0])
            sha = hashlib.sha1(blob["image"]).hexdigest()[:8]
            if sha not in bilder:
                ziel = os.path.join(bilderordner, f"{sha}.{blob['ext']}")
                with open(ziel, "wb") as f:
                    f.write(blob["image"])
                bilder[sha] = {"datei": os.path.basename(ziel), "pixel": None, "folien": []}
            bilder[sha]["folien"].append(nr)
            formen.append({"name": "Grafik", "typ": "PICTURE", "top": 1, "left": 0,
                           "platzhalter": None, "alt": "", "bild": sha})
        einheiten.append({"nr": nr, "layout": "pdf", "formen": formen, "notiz": ""})
    return einheiten, bilder


# --------------------------------------------------------------------- Bericht

def rumpf(einheit):
    """Alle Textabsätze einer Einheit ohne den Titel."""
    alle = [a["text"] for f in einheit["formen"] for a in f.get("absaetze", [])]
    return alle[1:] if alle else []


def bericht(einheiten, bilder):
    print(f"Einheiten: {len(einheiten)}")
    zeichen = sum(len(a["text"]) for e in einheiten
                  for f in e["formen"] for a in f.get("absaetze", []))
    print(f"Zeichen Text gesamt: {zeichen}")

    ohne_titel, textfelder, tabellen = [], 0, 0
    for e in einheiten:
        hat_titel = any(f["platzhalter"] and "TITLE" in f["platzhalter"]
                        and f.get("absaetze") for f in e["formen"])
        if not hat_titel:
            ohne_titel.append(e["nr"])
        for f in e["formen"]:
            if "TEXT_BOX" in f["typ"]:
                textfelder += 1
            if f.get("tabelle"):
                tabellen += 1

    print(f"Bilder: {sum(len(b['folien']) for b in bilder.values())} "
          f"Vorkommen, {len(bilder)} verschieden")
    ohne_alt = []
    for e in einheiten:
        for f in e["formen"]:
            if f.get("bild") and not brauchbarer_alt(f["alt"]):
                ohne_alt.append((e["nr"], f["bild"]))
    print(f"  ohne brauchbaren Alternativtext: {len(ohne_alt)}")
    mehrfach = [(s, len(b["folien"])) for s, b in bilder.items() if len(b["folien"]) > 2]
    if mehrfach:
        print("  mehrfach verwendet (Verdacht auf Dekoration): "
              + ", ".join(f"{s} ({n}x)" for s, n in mehrfach))
    print(f"Freie Textfelder: {textfelder} | Tabellen: {tabellen}")
    if ohne_titel:
        print(f"Einheiten ohne Titel: {ohne_titel}")

    print("\n--- Indizierte Absatzliste ---")
    for e in einheiten:
        alle = [a["text"] for f in e["formen"] for a in f.get("absaetze", [])]
        titel = alle[0] if alle else ""
        gfx = [f["bild"] for f in e["formen"] if f.get("bild")]
        print(f"\nF{e['nr']}: T={titel[:60]!r}" + (f"  Bilder={gfx}" if gfx else ""))
        for i, t in enumerate(rumpf(e)):
            print(f"   {i}: {t[:70]}")
        for f in e["formen"]:
            if f.get("tabelle"):
                print(f"   [TABELLE {len(f['tabelle'][0])} Sp. x {len(f['tabelle'])} Zei.]")


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("quelle")
    ap.add_argument("--out", default="arbeit", help="Arbeitsordner (Standard: arbeit)")
    args = ap.parse_args()

    bilderordner = os.path.join(args.out, "bilder")
    os.makedirs(bilderordner, exist_ok=True)

    endung = os.path.splitext(args.quelle)[1].lower()
    if endung == ".pptx":
        einheiten, bilder = aus_pptx(args.quelle, bilderordner)
    elif endung == ".docx":
        einheiten, bilder = aus_docx(args.quelle, bilderordner)
    elif endung == ".pdf":
        einheiten, bilder = aus_pdf(args.quelle, bilderordner)
    else:
        sys.exit(f"Nicht unterstützte Dateiendung: {endung}")

    ziel = os.path.join(args.out, "inhalt.json")
    with open(ziel, "w", encoding="utf-8") as f:
        json.dump({"quelle": os.path.basename(args.quelle),
                   "einheiten": einheiten, "bilder": bilder},
                  f, ensure_ascii=False, indent=1)

    bericht(einheiten, bilder)
    print(f"\ngeschrieben: {ziel}")
    print(f"Bilder in:   {bilderordner}  (jedes einzeln ansehen!)")


if __name__ == "__main__":
    main()
